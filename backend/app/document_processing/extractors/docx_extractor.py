import io
import os
import re
import hashlib
from typing import Optional, Callable, List, Dict, Any, Union
from .base_extractor import BaseDocumentExtractor
from ..models import (
    ExtractedDocument,
    DocumentSection,
    DocumentBlock,
    ExtractedTable,
    SourceLocation,
)
from ...utils.logger import setup_logger

logger = setup_logger(__name__)


class DocxExtractor(BaseDocumentExtractor):
    """
    High-performance DOCX financial report extractor:
    - Iterates body elements in true visual document sequence (interleaving paragraphs and tables)
    - Preserves heading hierarchy and section context
    - Handles merged cells, empty paragraphs, and multiple financial tables safely
    - Extracts currency symbols, units, fiscal years, and company name
    - Streams directly from file path or bytes without memory duplication
    """

    def extract(
        self,
        file_bytes: Union[bytes, str],
        document_id: str,
        filename: str,
        user_id: Optional[str] = None,
        company_name: Optional[str] = None,
        financial_year: Optional[str] = None,
        on_progress: Optional[Callable[[str, int, str], None]] = None,
    ) -> ExtractedDocument:
        import docx
        from docx.text.paragraph import Paragraph
        from docx.table import Table

        is_path = isinstance(file_bytes, str) and os.path.isfile(file_bytes)

        # 1. Compute Hash & File Size
        if is_path:
            file_path = str(file_bytes)
            file_size = os.path.getsize(file_path)
            hasher = hashlib.sha256()
            with open(file_path, "rb") as f_hash:
                for chunk in iter(lambda: f_hash.read(65536), b""):
                    hasher.update(chunk)
            doc_hash = hasher.hexdigest()
            doc = docx.Document(file_path)
        else:
            raw_bytes = file_bytes if isinstance(file_bytes, bytes) else str(file_bytes).encode("utf-8")
            file_size = len(raw_bytes)
            doc_hash = hashlib.sha256(raw_bytes).hexdigest()
            doc = docx.Document(io.BytesIO(raw_bytes))

        if on_progress:
            on_progress("Reading Word Document", 25, f"Parsing document structure, headings, and tables in {filename}")

        current_section = "Executive Summary"
        sections_map: Dict[str, List[DocumentBlock]] = {current_section: []}
        sections_tables: Dict[str, List[ExtractedTable]] = {current_section: []}

        detected_company = company_name
        detected_year = financial_year
        detected_currency = None
        detected_units = None

        filename_text = os.path.splitext(filename)[0].replace("_", " ")

        table_counter = 0

        # 2. Iterate Body Elements in Exact Visual Order
        for child in doc.element.body:
            tag = child.tag.lower()

            # A. Paragraphs
            if tag.endswith("p"):
                p = Paragraph(child, doc)
                text = p.text.strip()
                if not text:
                    continue

                style_name = p.style.name.lower() if p.style and p.style.name else ""
                is_heading = "heading" in style_name or "title" in style_name

                if is_heading:
                    current_section = text
                    if current_section not in sections_map:
                        sections_map[current_section] = []
                        sections_tables[current_section] = []

                # Detect metadata from paragraph text
                if not detected_company:
                    m_comp = re.search(r"([A-Z][A-Za-z0-9\s,&.-]+(?:Limited|Ltd|Corporation|Corp|Inc|LLC|Pvt|Bank|Manufacturing))", text, re.IGNORECASE)
                    if m_comp:
                        detected_company = m_comp.group(1).strip()
                if not detected_year:
                    m_fy = re.search(r"(FY\s*20\d\d|20\d\d\s*-\s*20\d\d|20\d\d)", text, re.IGNORECASE)
                    if m_fy:
                        detected_year = m_fy.group(0).strip()
                if not detected_currency:
                    m_curr = re.search(r"(\$|USD|INR|₹|EUR|€|GBP|£|Rs\.?)", text, re.IGNORECASE)
                    if m_curr:
                        detected_currency = m_curr.group(1).strip()
                if not detected_units:
                    m_unit = re.search(r"(in\s+crores?|in\s+lakhs?|in\s+millions?|in\s+billions?|in\s+thousands?)", text, re.IGNORECASE)
                    if m_unit:
                        detected_units = m_unit.group(1).strip()

                loc = SourceLocation(
                    format_type="docx",
                    section_name=current_section,
                    human_label=f'Section "{current_section}"'
                )

                sections_map[current_section].append(DocumentBlock(
                    block_type="heading" if is_heading else "text",
                    content=text,
                    source_location=loc,
                    section=current_section
                ))

            # B. Tables (Interleaved in Visual Order)
            elif tag.endswith("tbl"):
                table_counter += 1
                tbl = Table(child, doc)
                table_rows: List[List[str]] = []

                for row in tbl.rows:
                    row_cells = []
                    prev_cell_tc = None
                    for cell in row.cells:
                        # De-duplicate horizontal merged cells in python-docx
                        if prev_cell_tc is not None and getattr(cell, "_tc", None) == prev_cell_tc:
                            continue
                        c_text = cell.text.strip().replace("\n", " ")
                        row_cells.append(c_text)
                        prev_cell_tc = getattr(cell, "_tc", None)

                    if any(row_cells):
                        table_rows.append(row_cells)

                if not table_rows:
                    continue

                headers = table_rows[0]
                data_rows = table_rows[1:] if len(table_rows) > 1 else []

                loc = SourceLocation(
                    format_type="docx",
                    section_name=current_section,
                    human_label=f'Section "{current_section}", Table {table_counter}'
                )

                ext_table = ExtractedTable(
                    table_id=f"tbl_docx_{table_counter}",
                    headers=headers,
                    rows=data_rows,
                    title=f"Table {table_counter} ({current_section})",
                    source_location=loc,
                    units=detected_units,
                    currency=detected_currency,
                    table_type="financial_table"
                )
                tbl_md = ext_table.to_markdown()

                if current_section not in sections_map:
                    sections_map[current_section] = []
                    sections_tables[current_section] = []

                sections_tables[current_section].append(ext_table)
                sections_map[current_section].append(DocumentBlock(
                    block_type="table",
                    content=tbl_md,
                    table=ext_table,
                    source_location=loc,
                    section=current_section,
                    metadata={"table_index": table_counter, "headers": headers}
                ))

        if not detected_company:
            detected_company = filename_text.title()
        if not detected_year:
            detected_year = "FY2026"

        # 3. Build DocumentSection list
        doc_sections: List[DocumentSection] = []
        for sec_name, blks in sections_map.items():
            if not blks:
                continue
            loc = SourceLocation(
                format_type="docx",
                section_name=sec_name,
                human_label=f'Section "{sec_name}"'
            )
            raw_text = "\n\n".join([b.content for b in blks])
            doc_sections.append(DocumentSection(
                section_name=sec_name,
                source_location=loc,
                blocks=blks,
                raw_text=raw_text,
                tables=sections_tables.get(sec_name, []),
                metadata={"blocks_count": len(blks)}
            ))

        return ExtractedDocument(
            document_id=document_id,
            file_name=filename,
            file_type="docx",
            file_size=file_size,
            doc_hash=doc_hash,
            company_name=detected_company,
            financial_year=detected_year,
            currency=detected_currency,
            units=detected_units,
            sections=doc_sections,
            pages_or_sheets_count=len(doc_sections),
            metadata={
                "sections_count": len(doc_sections),
                "total_tables": table_counter,
            }
        )
