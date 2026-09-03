import os
import sys
import io
import json
import pytest

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))

from app.document_processing.file_detector import FileTypeDetector, UnsupportedFileFormatException
from app.document_processing.processor import DocumentProcessor
from app.document_processing.chunker import DocumentChunker
from app.vectorstore.vector_service import VectorStoreService


@pytest.fixture
def processor():
    return DocumentProcessor()


@pytest.fixture
def vector_service():
    return VectorStoreService(collection_name="test_multi_format_chunks")


def test_file_detector_formats():
    """Test format detection and validation across magic bytes and signatures."""
    # 1. PDF
    assert FileTypeDetector.detect_format(b"%PDF-1.7 mock pdf content", "report.pdf") == "pdf"
    
    # 2. CSV
    csv_bytes = b"Quarter,Revenue,Operating_Income,Net_Profit\nQ1 2026,12500,3200,2400\nQ2 2026,13100,3450,2600"
    assert FileTypeDetector.detect_format(csv_bytes, "financials.csv") == "csv"

    # 3. JSON
    json_bytes = b'{"company": "FinSight Tech", "financials": {"revenue_usd_m": 450.5, "ebitda_usd_m": 120.2}}'
    assert FileTypeDetector.detect_format(json_bytes, "metrics.json") == "json"

    # 4. Markdown
    md_bytes = b"# Q3 Financial Performance\n\n| Metric | Q3 2026 | Q3 2025 |\n| --- | --- | --- |\n| Revenue | $50M | $42M |"
    assert FileTypeDetector.detect_format(md_bytes, "quarterly_summary.md") == "md"

    # 5. Plain Text
    txt_bytes = b"Financial Notes for FY2026.\nRevenue grew by 15% year-over-year."
    assert FileTypeDetector.detect_format(txt_bytes, "notes.txt") == "txt"

    # 6. Unsupported format error
    with pytest.raises(UnsupportedFileFormatException):
        FileTypeDetector.detect_format(b"\x00\x00\x00\x18ftypmp42", "recording.mp4")

    with pytest.raises(UnsupportedFileFormatException):
        FileTypeDetector.validate_file(b"", "empty.pdf")


def test_csv_extraction_and_chunking(processor):
    """Test CSV structured parsing, markdown table generation, and chunk locations."""
    csv_content = """Fiscal Quarter,Total Revenue ($M),Net Income ($M),Operating Margin
Q1 FY26,4500,1100,24.4%
Q2 FY26,4750,1190,25.1%
Q3 FY26,4920,1240,25.2%
Q4 FY26,5200,1350,26.0%"""

    result = processor.process_document(
        file_source=csv_content.encode("utf-8"),
        document_id="doc_csv_test_1",
        user_id="test_user",
        file_name="quarterly_performance.csv",
        company_name="Acme Corp",
        financial_year="FY2026"
    )

    assert result.file_type == "csv"
    assert len(result.chunks) >= 1
    chunk = result.chunks[0]
    assert chunk.metadata["file_type"] == "csv"
    assert "CSV Rows" in chunk.metadata["source_location"]
    assert "Q1 FY26" in chunk.text
    assert "|" in chunk.text  # Markdown table verified


def test_json_extraction_and_chunking(processor):
    """Test JSON financial hierarchy parsing into key-value tables and array tables."""
    json_data = {
        "company_name": "Global Retailers Ltd",
        "financial_year": "FY2026",
        "currency": "USD",
        "balance_sheet": {
            "cash_and_equivalents": 125000000,
            "accounts_receivable": 45000000,
            "total_assets": 850000000,
            "total_liabilities": 320000000
        },
        "quarterly_breakdown": [
            {"quarter": "Q1", "revenue": 210000000, "net_income": 42000000},
            {"quarter": "Q2", "revenue": 225000000, "net_income": 47000000}
        ]
    }

    result = processor.process_document(
        file_source=json.dumps(json_data).encode("utf-8"),
        document_id="doc_json_test_1",
        user_id="test_user",
        file_name="financial_model.json"
    )

    assert result.file_type == "json"
    assert result.company_name == "Global Retailers Ltd"
    assert len(result.chunks) >= 1
    
    combined_chunk_text = " ".join([c.text for c in result.chunks])
    assert "cash_and_equivalents" in combined_chunk_text or "125000000" in combined_chunk_text
    assert any("JSON" in c.metadata["source_location"] for c in result.chunks)


def test_markdown_extraction_and_chunking(processor):
    """Test Markdown headings and Markdown table preservation."""
    md_content = """# Executive Financial Summary

Acme International reported record growth for FY2026.

## Revenue Breakdown

| Segment | FY2026 ($M) | FY2025 ($M) | Growth |
| --- | --- | --- | --- |
| Cloud Services | 1850 | 1200 | +54.2% |
| Enterprise Software | 920 | 850 | +8.2% |
| Consumer Products | 430 | 410 | +4.9% |

## Risk Factors

Supply chain inflation and currency volatility remain headwinds."""

    result = processor.process_document(
        file_source=md_content.encode("utf-8"),
        document_id="doc_md_test_1",
        user_id="test_user",
        file_name="annual_overview.md"
    )

    assert result.file_type == "md"
    assert len(result.chunks) >= 2
    
    # Check section detection
    sections = [c.section for c in result.chunks]
    assert any("Revenue Breakdown" in s or "Executive Financial Summary" in s for s in sections)
    
    # Check table integrity
    table_chunks = [c for c in result.chunks if c.metadata.get("is_table")]
    assert len(table_chunks) >= 1
    assert "Cloud Services" in table_chunks[0].text


def test_docx_extraction_and_chunking(processor):
    """Test Word document (.docx) extraction with paragraphs, headings, and Word tables."""
    import docx
    doc = docx.Document()
    doc.add_heading("Tata Consultancy Services Limited", level=0)
    doc.add_heading("Management Discussion & Analysis", level=1)
    doc.add_paragraph("Consolidated revenue for FY2026 reached new milestones across all digital verticals.")
    
    # Add a table
    table = doc.add_table(rows=3, cols=3)
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Vertical"
    hdr_cells[1].text = "Revenue (₹ Cr)"
    hdr_cells[2].text = "YoY Growth"

    row1 = table.rows[1].cells
    row1[0].text = "Banking & Financial Services"
    row1[1].text = "85400"
    row1[2].text = "14.2%"

    row2 = table.rows[2].cells
    row2[0].text = "Manufacturing & Retail"
    row2[1].text = "42100"
    row2[2].text = "11.8%"

    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    result = processor.process_document(
        file_source=docx_bytes,
        document_id="doc_docx_test_1",
        user_id="test_user",
        file_name="TCS_Annual_Report.docx"
    )

    assert result.file_type == "docx"
    assert result.company_name == "Tata Consultancy Services Limited"
    assert len(result.chunks) >= 1
    
    # Verify exact DOCX source location
    locs = [c.metadata.get("source_location", "") for c in result.chunks]
    assert any("Section" in loc for loc in locs)


def test_excel_xlsx_extraction_and_chunking(processor):
    """Test Excel (.xlsx) extraction preserving sheet names, row numbers, and Markdown tables."""
    import openpyxl
    wb = openpyxl.Workbook()
    
    # Sheet 1: Income Statement
    ws1 = wb.active
    ws1.title = "Income Statement"
    ws1.append(["Infosys Limited - Consolidated Income Statement (in USD Millions)", "", "", ""])
    ws1.append(["Metric", "FY2026", "FY2025", "YoY Change"])
    ws1.append(["Total Revenue", "19500", "18200", "+7.1%"])
    ws1.append(["Operating Profit", "4200", "3850", "+9.1%"])
    ws1.append(["Net Profit", "3250", "2980", "+9.1%"])

    # Sheet 2: Balance Sheet
    ws2 = wb.create_sheet(title="Balance Sheet")
    ws2.append(["Balance Sheet Item", "FY2026 ($M)", "FY2025 ($M)"])
    ws2.append(["Cash & Equivalents", "3100", "2800"])
    ws2.append(["Total Equity", "10200", "9400"])

    bio = io.BytesIO()
    wb.save(bio)
    xlsx_bytes = bio.getvalue()

    result = processor.process_document(
        file_source=xlsx_bytes,
        document_id="doc_xlsx_test_1",
        user_id="test_user",
        file_name="Infosys_Financials_FY2026.xlsx"
    )

    assert result.file_type == "xlsx"
    assert result.company_name == "Infosys Limited"
    assert len(result.chunks) >= 2
    
    # Check source location formatting: Sheet 'Income Statement', rows 2–5
    locs = [c.metadata.get("source_location", "") for c in result.chunks]
    assert any('Sheet "Income Statement"' in loc for loc in locs)
    assert any('Sheet "Balance Sheet"' in loc for loc in locs)


def test_end_to_end_multiformat_chromadb_and_rag(processor, vector_service):
    """Test indexing chunks from different formats and retrieving them with precise source locations."""
    # Process CSV
    csv_content = "Metric,FY2026\nCloud Revenue,1500\nOn-Premises,800"
    res_csv = processor.process_document(
        file_source=csv_content.encode("utf-8"),
        document_id="doc_e2e_csv",
        user_id="user_rag_test",
        file_name="cloud_metrics.csv"
    )
    
    # Process Markdown
    md_content = "# Operating Expenses\n\nTotal R&D expenditures reached $350M in FY2026."
    res_md = processor.process_document(
        file_source=md_content.encode("utf-8"),
        document_id="doc_e2e_md",
        user_id="user_rag_test",
        file_name="expenses.md"
    )

    # Upsert both into ChromaDB
    vector_service.add_document("doc_e2e_csv", "user_rag_test", res_csv.chunks, overwrite_if_exists=True)
    vector_service.add_document("doc_e2e_md", "user_rag_test", res_md.chunks, overwrite_if_exists=True)

    # Search for Cloud Revenue
    results_csv = vector_service.search("Cloud Revenue", document_id="doc_e2e_csv", user_id="user_rag_test", top_k=2)
    assert len(results_csv) >= 1
    assert "Cloud Revenue" in results_csv[0]["text"]
    assert "CSV" in (results_csv[0].get("source_location") or "")

    # Search for R&D expenditures
    results_md = vector_service.search("R&D expenditures", document_id="doc_e2e_md", user_id="user_rag_test", top_k=2)
    assert len(results_md) >= 1
    assert "350M" in results_md[0]["text"]
    assert "Operating Expenses" in (results_md[0].get("section") or "")
